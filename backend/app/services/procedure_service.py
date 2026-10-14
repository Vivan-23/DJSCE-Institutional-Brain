"""
Procedure (SOP) governance on top of skill extraction.

  - approve(): freezes a version (SkillVersion), sets owner office and a review date, and
    publishes the procedure into the evidence store so the assistant and ticket drafts can
    cite it ("Approved procedure: …"). Staff get the full text; for student-facing
    procedures, students get only the plain student summary.
  - edits to an approved procedure create a new pending version (back to draft, version+1);
    the approved copy stays searchable until the new version is approved.
  - retire / reject (reject needs a reason); stale procedures keep the reason and evidence.
  - feedback: repeated "wrong" reports flag a procedure for review.
  - export: one procedure or a handbook of every approved one, as DOCX or PDF.
"""
import io
import json
from datetime import datetime, timedelta, timezone
from typing import Any, Optional

import structlog
from sqlmodel import Session, select

from app.core.institution_config import ALL_DEPARTMENTS, get_institution_config
from app.models.structured import AuditLog, Skill, SkillVersion
from app.services.audit_service import audit_service
from app.services.fragment_service import fragment_service
from app.services.llm_service import llm_service

log = structlog.get_logger()

REVIEW_PERIOD_DAYS = 365
WRONG_REPORTS_TO_FLAG = 3
_LIST_FIELDS = ("preconditions", "steps", "decision_rules", "guardrails")
_TEXT_FIELDS = ("name", "description", "trigger", "escalation")


class ProcedureError(ValueError):
    """A rejected action; the message is safe to show."""


def _list(raw: Optional[str]) -> list[str]:
    try:
        return [str(x) for x in json.loads(raw or "[]")]
    except (json.JSONDecodeError, TypeError):
        return []


def snapshot(skill: Skill) -> dict[str, Any]:
    out: dict[str, Any] = {k: getattr(skill, k) for k in _TEXT_FIELDS}
    out.update({k: _list(getattr(skill, k)) for k in _LIST_FIELDS})
    out.update({"owner_office": skill.owner_office, "visibility": skill.visibility, "department": skill.department,
                "category": skill.category, "student_summary": skill.student_summary})
    return out


def full_text(skill: Skill) -> str:
    """The staff-facing searchable text of a procedure."""
    parts = [f"Approved procedure: {skill.name}. {skill.description or ''}", f"When it applies: {skill.trigger or '-'}"]
    for title, key in (("Before you start", "preconditions"), ("Steps", "steps"),
                       ("Decision rules", "decision_rules"), ("Must never happen", "guardrails")):
        items = _list(getattr(skill, key))
        if items:
            parts.append(f"{title}: " + " ".join(f"({i}) {s}" for i, s in enumerate(items, 1)))
    if skill.escalation:
        parts.append(f"Escalation: {skill.escalation}")
    if skill.owner_office:
        parts.append(f"Owned by: {skill.owner_office}")
    return "\n".join(parts)


class ProcedureService:
    # ── lifecycle ───────────────────────────────────────────────────────────
    def approve(self, db: Session, skill: Skill, actor: str, *, owner_office: Optional[str] = None,
                review_by: Optional[datetime] = None) -> Skill:
        if skill.status not in ("draft", "stale"):
            raise ProcedureError(f"Only a draft or a procedure needing review can be approved (this one is {skill.status})")
        if not _list(skill.steps):
            raise ProcedureError("A procedure needs at least one step before it can be approved")
        if skill.visibility == "public" and not (skill.student_summary or "").strip():
            skill.student_summary = self.draft_student_summary(skill)
        now = datetime.now(timezone.utc)
        skill.status = "approved"
        skill.owner_office = (owner_office or skill.owner_office or "").strip() or None
        skill.review_by = review_by or now + timedelta(days=REVIEW_PERIOD_DAYS)
        skill.approved_by = actor
        skill.approved_at = now
        skill.last_verified_at = now
        skill.stale_reason = None
        skill.stale_evidence = "[]"
        skill.reject_reason = None
        skill.updated_at = now
        db.add(skill)
        db.add(SkillVersion(skill_id=skill.id, version=skill.version, snapshot=json.dumps(snapshot(skill)),
                            approved_by=actor, approved_at=now))
        db.commit()
        db.refresh(skill)
        self.publish(db, skill)
        return skill

    def publish(self, db: Session, skill: Skill) -> int:
        """(Re)index the approved procedure for retrieval; returns the number of excerpts written."""
        fragment_service.delete_for_skill(db, skill.id)
        if skill.status != "approved":
            return 0
        label = f"Approved procedure: {skill.name} (v{skill.version})"
        fragment_service.create_fragment(db, source_type="procedure", source_ref=label, content=full_text(skill),
                                         visibility="internal", department=skill.department, skill_id=skill.id)
        written = 1
        if skill.visibility == "public" and (skill.student_summary or "").strip():
            fragment_service.create_fragment(
                db, source_type="procedure", source_ref=label,
                content=f"{skill.name} (what students need to do): {skill.student_summary}",
                visibility="public", department=skill.department, skill_id=skill.id,
            )
            written += 1
        return written

    def publish_missing(self, db: Session) -> int:
        """Index approved procedures that have no searchable copy yet (e.g. approved before this existed)."""
        from app.models.structured import Fragment

        done = 0
        for skill in db.exec(select(Skill).where(Skill.status == "approved")).all():
            if db.exec(select(Fragment).where(Fragment.skill_id == skill.id)).first() is None:
                self.publish(db, skill)
                done += 1
        if done:
            log.info("Published approved procedures", count=done)
        return done

    def start_revision(self, db: Session, skill: Skill) -> Skill:
        """Editing an approved procedure opens a new pending version; the approved copy stays live."""
        if skill.status == "approved":
            skill.version += 1
            skill.status = "draft"
            db.add(skill)
            db.commit()
            db.refresh(skill)
        return skill

    def reject(self, db: Session, skill: Skill, actor: str, reason: str) -> Skill:
        if not reason.strip():
            raise ProcedureError("Give a reason for rejecting")
        if skill.status not in ("draft", "stale"):
            raise ProcedureError("Only a draft or a procedure needing review can be rejected")
        skill.status = "rejected"
        skill.reject_reason = reason.strip()
        skill.updated_at = datetime.now(timezone.utc)
        db.add(skill)
        db.commit()
        if not self._has_approved_version(db, skill):  # never approved: nothing should stay searchable
            fragment_service.delete_for_skill(db, skill.id)
        db.refresh(skill)
        return skill

    def retire(self, db: Session, skill: Skill, actor: str) -> Skill:
        skill.status = "deprecated"
        skill.updated_at = datetime.now(timezone.utc)
        db.add(skill)
        db.commit()
        fragment_service.delete_for_skill(db, skill.id)
        db.refresh(skill)
        return skill

    @staticmethod
    def _has_approved_version(db: Session, skill: Skill) -> bool:
        return db.exec(select(SkillVersion).where(SkillVersion.skill_id == skill.id)).first() is not None

    def create_manual(self, db: Session, fields: dict[str, Any], actor: str) -> Skill:
        name = (fields.get("name") or "").strip()
        if not name:
            raise ProcedureError("A procedure needs a name")
        skill = Skill(
            name=name, description=(fields.get("description") or "").strip(), trigger=(fields.get("trigger") or "").strip(),
            escalation=(fields.get("escalation") or "").strip() or None, status="draft", origin="manual",
            **{k: json.dumps([s.strip() for s in fields.get(k) or [] if str(s).strip()]) for k in _LIST_FIELDS},
        )
        self.apply_settings(skill, fields)
        db.add(skill)
        db.commit()
        db.refresh(skill)
        audit_service.record(db, actor=actor, action="skill_create", target_id=str(skill.id))
        return skill

    @staticmethod
    def apply_settings(skill: Skill, fields: dict[str, Any]) -> None:
        """Owner office, visibility, department, category, student summary, review date (validated)."""
        cfg = get_institution_config()
        if "visibility" in fields and fields["visibility"] is not None:
            if fields["visibility"] not in ("internal", "public"):
                raise ProcedureError("Visibility must be internal or public")
            skill.visibility = fields["visibility"]
        if fields.get("department") is not None:
            dep = str(fields["department"]).strip()
            if dep.lower() == ALL_DEPARTMENTS or not dep:
                skill.department = ALL_DEPARTMENTS
            else:
                code = cfg.department_code(dep)
                if not code:
                    raise ProcedureError(f"Unknown department '{dep}'")
                skill.department = code
        if fields.get("category") is not None:
            cat = str(fields["category"]).strip().lower() or None
            if cat and cat not in {c.key for c in cfg.helpdesk_categories}:
                raise ProcedureError(f"Unknown category '{cat}'")
            skill.category = cat
        for key in ("owner_office", "student_summary"):
            if fields.get(key) is not None:
                setattr(skill, key, str(fields[key]).strip() or None)
        if fields.get("review_by") is not None:
            skill.review_by = fields["review_by"]

    def draft_student_summary(self, skill: Skill) -> Optional[str]:
        """A plain 'what you need to do' paragraph for students, without internal steps or guardrails."""
        if not llm_service.is_available:
            return None
        text = llm_service.complete(
            "You rewrite an internal college procedure as 2-4 plain sentences telling a STUDENT what they need to "
            "do, by when, and whom to contact. Leave out internal staff steps, decision rules and guardrails. "
            "Use only facts in the procedure; never invent dates or fees.",
            full_text(skill), max_tokens=300, temperature=0.2,
        )
        return (text or "").strip() or None

    # ── reading ─────────────────────────────────────────────────────────────
    def versions(self, db: Session, skill: Skill) -> list[dict[str, Any]]:
        rows = db.exec(select(SkillVersion).where(SkillVersion.skill_id == skill.id).order_by(SkillVersion.version.desc()))
        return [{"version": v.version, "approved_by": v.approved_by, "approved_at": v.approved_at.isoformat(),
                 "snapshot": json.loads(v.snapshot)} for v in rows]

    def feedback_counts(self, db: Session, skill: Skill) -> dict[str, int]:
        counts = {"helpful": 0, "wrong": 0}
        for e in db.exec(select(AuditLog).where(AuditLog.action == "skill_use", AuditLog.target_id == str(skill.id))):
            outcome = json.loads(e.audit_metadata or "{}").get("outcome")
            if outcome in ("success", "helpful"):
                counts["helpful"] += 1
            elif outcome in ("failure", "wrong"):
                counts["wrong"] += 1
        return counts

    def after_feedback(self, db: Session, skill: Skill) -> bool:
        """Flag an approved procedure for review once enough people report it as wrong."""
        if skill.status != "approved" or self.feedback_counts(db, skill)["wrong"] < WRONG_REPORTS_TO_FLAG:
            return False
        skill.status = "stale"
        skill.stale_reason = f"{WRONG_REPORTS_TO_FLAG} or more people reported this procedure as wrong or outdated."
        db.add(skill)
        db.commit()
        return True

    def record_use(self, db: Session, skill_ids: set[str]) -> None:
        """Count how often approved procedures back an answer or a ticket draft."""
        from uuid import UUID

        for sid in skill_ids:
            try:
                skill = db.get(Skill, UUID(sid))
            except ValueError:
                continue
            if skill:
                skill.use_count += 1
                db.add(skill)
        db.commit()

    # ── export ──────────────────────────────────────────────────────────────
    def _sections(self, skill: Skill) -> list[tuple[str, list[str], bool]]:
        return [("Before you start", _list(skill.preconditions), False), ("Steps", _list(skill.steps), True),
                ("Decision rules", _list(skill.decision_rules), False),
                ("Must never happen", _list(skill.guardrails), False)]

    def _meta_line(self, skill: Skill) -> str:
        bits = [f"Version {skill.version}", f"status: {skill.status}"]
        if skill.owner_office:
            bits.append(f"owner: {skill.owner_office}")
        if skill.approved_by:
            bits.append(f"approved by {skill.approved_by} on {skill.approved_at.date().isoformat()}")
        if skill.review_by:
            bits.append(f"review by {skill.review_by.date().isoformat()}")
        return " · ".join(bits)

    def to_docx(self, skills: list[Skill], title: str) -> bytes:
        from docx import Document
        from docx.shared import Pt

        doc = Document()
        doc.add_paragraph(get_institution_config().full_name).runs[0].bold = True
        doc.add_heading(title, level=0 if len(skills) > 1 else 1)
        note = doc.add_paragraph("Demonstration content generated by the DJSCE Institutional Brain prototype; "
                                 "not official college policy.")
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(8)
        for skill in skills:
            doc.add_heading(skill.name, level=1 if len(skills) > 1 else 2)
            doc.add_paragraph(self._meta_line(skill)).runs[0].font.size = Pt(9)
            if skill.description:
                doc.add_paragraph(skill.description)
            doc.add_paragraph(f"When it applies: {skill.trigger or '-'}")
            for heading, items, numbered in self._sections(skill):
                if items:
                    doc.add_heading(heading, level=3)
                    for item in items:
                        doc.add_paragraph(item, style="List Number" if numbered else "List Bullet")
            if skill.escalation:
                doc.add_heading("Escalation", level=3)
                doc.add_paragraph(skill.escalation)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def to_pdf(self, skills: list[Skill], title: str) -> bytes:
        from xml.sax.saxutils import escape

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import ListFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer

        st = getSampleStyleSheet()
        small = st["BodyText"].clone("small", fontSize=8, leading=10, textColor=colors.grey)
        story = [Paragraph(escape(get_institution_config().full_name), st["Heading4"]), Paragraph(escape(title), st["Title"]),
                 Paragraph("<i>Demonstration content generated by the DJSCE Institutional Brain prototype; not official "
                           "college policy.</i>", small), Spacer(1, 0.4 * cm)]
        for n, skill in enumerate(skills):
            if n:
                story.append(PageBreak())
            story += [Paragraph(escape(skill.name), st["Heading1"]), Paragraph(escape(self._meta_line(skill)), small)]
            if skill.description:
                story.append(Paragraph(escape(skill.description), st["BodyText"]))
            story.append(Paragraph(f"<b>When it applies:</b> {escape(skill.trigger or '-')}", st["BodyText"]))
            for heading, items, numbered in self._sections(skill):
                if items:
                    story.append(Paragraph(escape(heading), st["Heading3"]))
                    story.append(ListFlowable([Paragraph(escape(i), st["BodyText"]) for i in items],
                                              bulletType="1" if numbered else "bullet"))
            if skill.escalation:
                story += [Paragraph("Escalation", st["Heading3"]), Paragraph(escape(skill.escalation), st["BodyText"])]
        buf = io.BytesIO()
        SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=2 * cm, rightMargin=2 * cm,
                          topMargin=1.8 * cm, bottomMargin=1.8 * cm).build(story)
        return buf.getvalue()


procedure_service = ProcedureService()
