"""Timetable file → transcribed boxes ``{day, start, end, text}`` + header + legend (no interpretation).

- .docx / .xlsx / .csv: the grid is read directly; merged cells become one box spanning several rows.
- images (.png/.jpg/.webp) and .pdf (first page rendered to an image): a vision model transcribes the sheet
  twice; boxes the two readings disagree on are marked ``check`` so the reviewer looks at them first.

Interpretation (subject, batch, faculty, room) happens in ``timetable_parser``; publication always needs a
person to review the draft (``timetable_service``).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import re
from typing import Any, Optional

import structlog

from app.core.config import settings
from app.services.llm_service import llm_service
from app.services.timetable_parser import normalize_time, weekday_of

log = structlog.get_logger()

IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp"}
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_TIME_RANGE = re.compile(r"(\d{1,2}[:.]\d{2})\s*(?:-|–|—|to)\s*(\d{1,2}[:.]\d{2})")


class ExtractionError(RuntimeError):
    pass


def kind_of(filename: str, content_type: str) -> Optional[str]:
    name = filename.lower()
    if name.endswith(".docx"):
        return "docx"
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".csv"):
        return "csv"
    if name.endswith(".pdf") or content_type == "application/pdf":
        return "pdf"
    if content_type in IMAGE_TYPES or re.search(r"\.(png|jpe?g|webp)$", name):
        return "image"
    return None


# ── grids (docx / xlsx / csv) ────────────────────────────────────────────────
# A grid is a matrix of (box_id, text); merged cells share a box_id.

def _grid_to_boxes(grid: list[list[tuple[Any, str]]]) -> tuple[list[dict], list[str]]:
    """Find the day header row and the time column, then emit one box per merged region."""
    notes: list[str] = []
    day_counts = [sum(weekday_of(t) is not None for _, t in row if t) for row in grid]
    best = max(day_counts, default=0)
    header_row = day_counts.index(best) if best else None
    if header_row is None:
        raise ExtractionError("No row with weekday names (Monday, Tuesday, …) was found in the table.")
    col_day = {c: weekday_of(t) for c, (_, t) in enumerate(grid[header_row]) if weekday_of(t) is not None}
    time_col = next((c for c in range(len(grid[header_row]))
                     if any(_TIME_RANGE.search(row[c][1] or "") for row in grid[header_row + 1:] if c < len(row))), None)
    if time_col is None:
        raise ExtractionError("No column with time ranges (e.g. 8:00-9:00) was found in the table.")
    row_time: dict[int, tuple[str, str]] = {}
    for r in range(header_row + 1, len(grid)):
        m = _TIME_RANGE.search(grid[r][time_col][1] or "") if time_col < len(grid[r]) else None
        if m:
            row_time[r] = (m.group(1).replace(".", ":"), m.group(2).replace(".", ":"))
    seen: dict[Any, dict] = {}
    last_in_col: dict[int, dict] = {}
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    for r in sorted(row_time):
        for c, (box_id, text) in enumerate(grid[r]):
            if c == time_col or c not in col_day or not (text or "").strip():
                last_in_col.pop(c, None)
                continue
            clean = re.sub(r"\s+", " ", text).strip()
            key = (box_id, col_day[c])
            prev = last_in_col.get(c)
            if key in seen:
                seen[key]["end"] = row_time[r][1]
            elif prev and prev["text"] == clean and prev["end"] == row_time[r][0]:
                # Same text directly below: a merged cell that the format (or library) split per row.
                prev["end"] = row_time[r][1]
                seen[key] = prev
            else:
                seen[key] = {"day": days[col_day[c]], "start": row_time[r][0], "end": row_time[r][1], "text": clean}
            last_in_col[c] = seen[key]
    seen = {id(v): v for v in seen.values()}
    if not seen:
        notes.append("The table had day and time headings but no filled cells.")
    return list(seen.values()), notes


def _header_from_text(lines: list[str]) -> dict[str, Optional[str]]:
    blob = " \n".join(lines)

    def grab(pattern: str) -> Optional[str]:
        m = re.search(pattern, blob, re.I)
        return m.group(0) if m else None

    return {
        "department": grab(r"Department of [A-Za-z &]+"),
        "class": grab(r"Class\s*:\s*[^\n|]+?(?=\s{2,}|Sem|Div|$)"),
        "semester": grab(r"Sem(ester)?\s*:\s*[IVX\d]+"),
        "division": grab(r"Div(ision)?\s*:\s*[A-Z]\d{1,2}"),
        "term": grab(r"Academic Term\s*:\s*[A-Za-z]+\s*-\s*[A-Za-z]+\s*\d{4}"),
        "effective_from": grab(r"w\.?\s*e\.?\s*f\.?\s*:?\s*[A-Za-z0-9 ,./-]+\d{4}"),
    }


def _from_docx(data: bytes) -> dict:
    import docx

    doc = docx.Document(io.BytesIO(data))
    text_lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        # Box identity by position (never id() of python-docx proxies, which Python may reuse): a run of identical
        # text across a row is one horizontally merged box; vertical continuation is handled in _grid_to_boxes.
        grid = []
        for r, row in enumerate(table.rows):
            out, start, prev = [], 0, None
            for c, cell in enumerate(row.cells):
                text = cell.text
                if c == 0 or text != prev or not text.strip():
                    start = c
                out.append(((r, start), text))
                prev = text
            grid.append(out)
        # Header text often sits in the first table rows too.
        text_lines += [" ".join(dict.fromkeys(t for _, t in row if t)) for row in grid[:4]]
        try:
            boxes, notes = _grid_to_boxes(grid)
        except ExtractionError:
            continue
        return {"cells": boxes, "notes": notes, "header": _header_from_text(text_lines), "legend": []}
    raise ExtractionError("No timetable grid (weekday header + time column) was found in the document.")


def _from_xlsx(data: bytes) -> dict:
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).active
    owner: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                owner[(r, c)] = (rng.min_row, rng.min_col)
    grid = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            o = owner.get((r, c), (r, c))
            v = ws.cell(*o).value
            row.append((o, "" if v is None else str(v)))
        grid.append(row)
    lines = [" ".join(dict.fromkeys(t for _, t in row if t)) for row in grid[:6]]
    boxes, notes = _grid_to_boxes(grid)
    return {"cells": boxes, "notes": notes, "header": _header_from_text(lines), "legend": []}


def _from_csv(data: bytes) -> dict:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # CSV has no merges: identical text directly below in the same column continues the box.
    grid, prev = [], {}
    for r, row in enumerate(rows):
        out = []
        for c, text in enumerate(row):
            box = prev.get(c) if text and prev.get(c, (None, None))[1] == text else ((r, c), text)
            prev[c] = box
            out.append((box[0], text))
        grid.append(out)
    lines = [" ".join(t for t in row if t) for row in rows[:6]]
    boxes, notes = _grid_to_boxes(grid)
    return {"cells": boxes, "notes": notes, "header": _header_from_text(lines), "legend": []}


# ── images / pdf (vision transcription) ─────────────────────────────────────
VISION_PROMPT = """You transcribe a college class timetable image into JSON. Copy text exactly as printed; do not
interpret, expand abbreviations, or guess. Return ONLY a JSON object:
{"header": {"department": str|null, "class": str|null, "semester": str|null, "division": str|null,
            "term": str|null, "effective_from": str|null},
 "cells": [{"day": "Monday".."Saturday", "start": "H:MM", "end": "H:MM", "text": str}],
 "legend": [{"code": str, "subject": str, "faculty": str}],
 "unreadable": [str]}
Rules:
- One entry in "cells" per printed box. Read each box's time from the time label printed at the start of its row.
  A box spanning several time rows gets the start of its first row and the end of its last row.
- When a day column is split into side-by-side boxes in the same time slot, output each box separately.
- Skip BREAK and empty boxes.
- "legend" copies the subject/faculty key table, one row per subject.
- If text is unreadable, describe it in "unreadable" instead of guessing."""


def _image_bytes(kind: str, data: bytes, content_type: str) -> tuple[bytes, str]:
    from PIL import Image

    if kind == "pdf":
        import fitz  # pymupdf

        doc = fitz.open(stream=data, filetype="pdf")
        if doc.page_count == 0:
            raise ExtractionError("The PDF has no pages.")
        data = doc[0].get_pixmap(dpi=200).tobytes("png")
        content_type = "image/png"
    img = Image.open(io.BytesIO(data))
    img.load()
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    longest = max(img.size)
    if longest > 2400 or len(data) > 3_500_000:
        scale = min(1.0, 2400 / longest)
        img = img.resize((int(img.width * scale), int(img.height * scale)))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, "JPEG", quality=88)
        return buf.getvalue(), "image/jpeg"
    return data, content_type if content_type in IMAGE_TYPES else "image/png"


def _transcribe(image: bytes, content_type: str) -> dict:
    url = f"data:{content_type};base64,{base64.b64encode(image).decode()}"
    completion = llm_service.client.chat.completions.create(
        model=settings.TIMETABLE_VISION_MODEL,
        messages=[{"role": "user", "content": [{"type": "text", "text": VISION_PROMPT},
                                               {"type": "image_url", "image_url": {"url": url}}]}],
        max_tokens=6000, temperature=0, response_format={"type": "json_object"},
    )
    return json.loads(completion.choices[0].message.content)


def _box_key(c: dict) -> tuple:
    return (weekday_of(c.get("day")), normalize_time(c.get("start")), normalize_time(c.get("end")),
            re.sub(r"[\s_]", "", str(c.get("text", ""))).upper())


def _from_image(kind: str, data: bytes, content_type: str) -> dict:
    if not llm_service.is_available:
        raise ExtractionError("Reading photos and PDFs needs the language model (GROQ_API_KEY). Upload a Word, Excel "
                              "or CSV timetable instead, or add the entries by hand on the review screen.")
    image, ctype = _image_bytes(kind, data, content_type)
    first = _transcribe(image, ctype)
    notes = [f"Unreadable: {u}" for u in first.get("unreadable") or []]
    cells = [c for c in first.get("cells") or [] if isinstance(c, dict)]
    try:
        second = _transcribe(image, ctype)
        other = {_box_key(c) for c in second.get("cells") or [] if isinstance(c, dict)}
        differing = 0
        for c in cells:
            if _box_key(c) not in other:
                c["check"] = True
                differing += 1
        if differing:
            notes.append(f"{differing} cell(s) read differently on a second pass; they are marked 'check against source'.")
    except Exception as e:  # the second reading is a cross-check, not a requirement
        log.warning("timetable second transcription failed", error=str(e))
        notes.append("Only one reading was possible; check every cell against the source.")
    return {"cells": cells, "notes": notes, "header": first.get("header") or {}, "legend": first.get("legend") or []}


def extract(filename: str, content_type: str, data: bytes) -> dict:
    """→ {method, cells, header, legend, notes}. Raises ExtractionError with a reviewer-readable message."""
    kind = kind_of(filename, content_type)
    if kind is None:
        raise ExtractionError("Unsupported file. Upload a photo (PNG/JPG), PDF, Word (.docx), Excel (.xlsx) or CSV.")
    if kind == "docx":
        out = _from_docx(data)
    elif kind == "xlsx":
        out = _from_xlsx(data)
    elif kind == "csv":
        out = _from_csv(data)
    else:
        out = _from_image(kind, data, content_type)
    out["method"] = "vision" if kind in ("image", "pdf") else kind
    return out
