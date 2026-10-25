"""Timetable upload → review → publish, with synthetic sheets (made-up initials and accounts)."""
import csv
import io
import json
from datetime import date, timedelta
from unittest.mock import MagicMock, patch
from uuid import UUID

import pytest
from sqlmodel import Session, select

from app.models.structured import ScheduleException, TimetableEntry, User
from app.services.timetable_service import timetable_service
from tests.helpers import register

HEADER = ["Department of Computer Engineering", "Class: S.Y. B.TECH.   Sem: III   Div: {div}   w.e.f. July 1, 2026"]


def _h(token):
    return {"Authorization": f"Bearer {token}"}


def _docx(div="Z1", rows=None) -> bytes:
    """Days across, times down; the 8-10 lab box is a vertically merged cell."""
    import docx

    rows = rows or [
        ["Time", "Monday", "Tuesday", "Wednesday"],
        ["8:00-9:00", "DSA_Z11 (QX) L1", "OOP (RY) 41", ""],
        ["9:00-10:00", "DSA_Z11 (QX) L1", "BREAK", "DBMS (QX) 42"],
        ["10:00-11:00", "MATH (QZ 41, RY 42)", "", ""],
    ]
    d = docx.Document()
    for line in HEADER:
        d.add_paragraph(line.format(div=div))
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    if rows[1][1] and rows[1][1] == rows[2][1]:
        t.cell(1, 1).merge(t.cell(2, 1)).text = rows[1][1]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture()
def run_sync(test_engine):
    """Run the extraction task inline instead of through Celery."""
    def _run(import_id):
        with Session(test_engine) as s:
            timetable_service.run_extraction(s, UUID(import_id))
    with patch("app.core.celery.process_timetable_task.delay", side_effect=_run):
        yield


def _upload(client, token, data, name="sy_z1.docx", ctype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
    return client.post("/api/v1/timetables", headers=_h(token), files={"file": (name, data, ctype)})


def _user_id(db, email):
    return str(db.exec(select(User).where(User.email == email)).first().id)


def _setup_people(client, prefix):
    coord = register(client, f"{prefix}-coord@test.com", role="manager")
    register(client, f"{prefix}-qx@test.com", full_name="Quinn Xavier", role="employee")
    register(client, f"{prefix}-ry@test.com", full_name="Riya Yadav", role="employee")
    register(client, f"{prefix}-qz@test.com", full_name="Quentin Zed", role="employee")
    return coord


def _map(client, db, coord, prefix):
    for abbr, who in (("QX", "qx"), ("RY", "ry"), ("QZ", "qz")):
        existing = client.get("/api/v1/timetables/aliases", headers=_h(coord)).json()
        for a in existing:
            if a["abbrev"] == abbr:
                client.delete(f"/api/v1/timetables/aliases/{a['id']}", headers=_h(coord))
        res = client.post("/api/v1/timetables/aliases", headers=_h(coord),
                          json={"abbrev": abbr, "user_id": _user_id(db, f"{prefix}-{who}@test.com")})
        assert res.status_code == 201, res.text


def _monday(weeks=1):
    t = date.today()
    return t - timedelta(days=t.weekday()) + timedelta(weeks=weeks)


def test_docx_upload_extracts_a_draft_that_cannot_publish_with_unknown_initials(client, db_session, run_sync):
    coord = _setup_people(client, "tt1")
    res = _upload(client, coord, _docx("Z1"))
    assert res.status_code == 202, res.text
    tid = res.json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    assert d["status"] == "draft" and d["class_key"] == "SY-Z1" and d["effective_from"] == "2026-07-01"
    lab = next(e for e in d["entries"] if e["subject"] == "DSA")
    assert (lab["day"], lab["start_time"], lab["end_time"], lab["batch"], lab["kind"]) == ("Monday", "08:00", "10:00", "Z11", "lab")
    math = sorted((e["faculty_abbrev"], e["room"]) for e in d["entries"] if e["subject"] == "MATH")
    assert math == [("QZ", "41"), ("RY", "42")]
    assert all(any(i["code"] == "unknown_faculty" for i in e["issues"]) for e in d["entries"])
    res = client.post(f"/api/v1/timetables/{tid}/publish", headers=_h(coord), json={"acknowledge_warnings": True})
    assert res.status_code == 400 and "need fixing" in res.json()["detail"]


def test_map_publish_and_schedule_keeps_source(client, db_session, run_sync):
    coord = _setup_people(client, "tt2")
    tid = _upload(client, coord, _docx("Z2"), name="sy_z2.docx").json()["id"]
    sugg = client.get(f"/api/v1/timetables/{tid}/suggestions", headers=_h(coord)).json()
    qx = next(s for s in sugg if s["abbrev"] == "QX")
    assert any(c["email"] == "tt2-qx@test.com" and "initials" in c["reason"] for c in qx["candidates"])
    _map(client, db_session, coord, "tt2")
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    assert d["counts"]["blocking"] == 0
    # Warnings (here: a missing room on MATH? no — all have rooms; DSA has rooms) still need confirmation if present.
    res = client.post(f"/api/v1/timetables/{tid}/publish", headers=_h(coord), json={"acknowledge_warnings": True})
    assert res.status_code == 200, res.text
    assert res.json()["class_key"] == "SY-Z2"
    fac = client.post("/api/v1/auth/login", json={"email": "tt2-qx@test.com", "password": "pass1234"}).json()["access_token"]
    sched = client.get(f"/api/v1/schedule/mine?start={_monday().isoformat()}&days=7", headers=_h(fac)).json()
    subjects = sorted((i["subject"], i["start_time"]) for i in sched["items"])
    assert subjects == [("DBMS", "09:00"), ("DSA", "08:00")]
    assert all(i["source"]["filename"] == "sy_z2.docx" for i in sched["items"])
    inbox = client.get("/api/v1/notices/mine", headers=_h(fac)).json()
    assert any("Timetable published" in n["title"] for n in inbox)
    # The source file is kept for review.
    src = client.get(f"/api/v1/timetables/{tid}/source", headers=_h(coord))
    assert src.status_code == 200 and len(src.content) > 1000


def test_same_file_twice_is_not_duplicated(client, db_session, run_sync):
    coord = _setup_people(client, "tt3")
    data = _docx("Z3")
    first = _upload(client, coord, data, name="z3.docx")
    again = _upload(client, coord, data, name="z3-copy.docx")
    assert first.status_code == 202 and again.status_code == 200
    assert again.json()["duplicate"] is True and again.json()["id"] == first.json()["id"]


def test_reupload_diff_supersedes_and_carries_exceptions(client, db_session, run_sync):
    coord = _setup_people(client, "tt4")
    _map(client, db_session, coord, "tt4")
    v1 = _upload(client, coord, _docx("Z4"), name="z4-v1.docx").json()["id"]
    assert client.post(f"/api/v1/timetables/{v1}/publish", headers=_h(coord), json={"acknowledge_warnings": True}).status_code == 200
    oop = db_session.exec(select(TimetableEntry).where(TimetableEntry.import_id == UUID(v1),
                                                       TimetableEntry.subject == "OOP")).first()
    tuesday = _monday() + timedelta(days=1)
    assert client.post("/api/v1/schedule/exceptions", headers=_h(coord), json={
        "on_date": tuesday.isoformat(), "kind": "cancelled", "entry_id": str(oop.id)}).status_code == 201

    rows = [
        ["Time", "Monday", "Tuesday", "Wednesday"],
        ["8:00-9:00", "DSA_Z11 (QX) L1", "OOP (RY) 43", ""],       # OOP room changed
        ["9:00-10:00", "DSA_Z11 (QX) L1", "BREAK", ""],             # DBMS removed
        ["10:00-11:00", "MATH (QZ 41, RY 42)", "", "CN (QZ) 44"],   # CN added
    ]
    v2 = _upload(client, coord, _docx("Z4", rows), name="z4-v2.docx").json()["id"]
    d = client.get(f"/api/v1/timetables/{v2}", headers=_h(coord)).json()
    assert d["version"] == 2
    assert [e["subject"] for e in d["diff"]["added"]] == ["CN"]
    assert [e["subject"] for e in d["diff"]["removed"]] == ["DBMS"]
    assert [(c["before"]["room"], c["after"]["room"]) for c in d["diff"]["changed"]] == [("41", "43")]
    res = client.post(f"/api/v1/timetables/{v2}/publish", headers=_h(coord), json={"acknowledge_warnings": True})
    assert res.status_code == 200 and res.json()["exceptions_carried_over"] == 1
    old = client.get(f"/api/v1/timetables/{v1}", headers=_h(coord)).json()
    assert old["status"] == "superseded"
    ry = client.post("/api/v1/auth/login", json={"email": "tt4-ry@test.com", "password": "pass1234"}).json()["access_token"]
    items = client.get(f"/api/v1/schedule/mine?start={tuesday.isoformat()}&days=1", headers=_h(ry)).json()["items"]
    oop_item = next(i for i in items if i["subject"] == "OOP")
    assert (oop_item["room"], oop_item["status"]) == ("43", "cancelled")


def test_two_classes_consolidate_and_clashes_are_flagged(client, db_session, run_sync):
    coord = _setup_people(client, "tt5")
    _map(client, db_session, coord, "tt5")
    a = _upload(client, coord, _docx("Z5"), name="z5.docx").json()["id"]
    assert client.post(f"/api/v1/timetables/{a}/publish", headers=_h(coord), json={"acknowledge_warnings": True}).status_code == 200
    rows = [["Time", "Monday", "Tuesday"], ["8:00-9:00", "", "ALGO (QX) 50"], ["9:00-10:00", "MATH (QX) 51", ""]]
    b = _upload(client, coord, _docx("Z6", rows), name="z6.docx").json()["id"]
    d = client.get(f"/api/v1/timetables/{b}", headers=_h(coord)).json()
    math = next(e for e in d["entries"] if e["subject"] == "MATH")
    assert any(i["code"] == "clash" for i in math["issues"])  # QX already teaches DSA Mon 8-10 in Z5
    res = client.post(f"/api/v1/timetables/{b}/publish", headers=_h(coord), json={"acknowledge_warnings": False})
    assert res.status_code == 400 and "warnings" in res.json()["detail"]
    client.delete(f"/api/v1/timetables/{b}/entries/{math['id']}", headers=_h(coord))
    assert client.post(f"/api/v1/timetables/{b}/publish", headers=_h(coord), json={"acknowledge_warnings": False}).status_code == 200
    qx = client.post("/api/v1/auth/login", json={"email": "tt5-qx@test.com", "password": "pass1234"}).json()["access_token"]
    items = client.get(f"/api/v1/schedule/mine?start={_monday().isoformat()}&days=7", headers=_h(qx)).json()["items"]
    assert {i["class_label"] for i in items} == {"SY Z5", "SY Z6"}


def test_reviewer_corrections(client, db_session, run_sync):
    coord = _setup_people(client, "tt6")
    _map(client, db_session, coord, "tt6")
    tid = _upload(client, coord, _docx("Z7"), name="z7.docx").json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    e = d["entries"][0]
    res = client.patch(f"/api/v1/timetables/{tid}/entries/{e['id']}", headers=_h(coord), json={"room": "99", "start_time": "7am"})
    assert res.status_code == 400
    res = client.patch(f"/api/v1/timetables/{tid}/entries/{e['id']}", headers=_h(coord), json={"room": "99"})
    assert res.status_code == 200 and res.json()["room"] == "99" and res.json()["edited"] is True
    added = client.post(f"/api/v1/timetables/{tid}/entries", headers=_h(coord), json={
        "weekday": 4, "start_time": "14:00", "end_time": "15:00", "subject": "SEM", "room": "45", "faculty_abbrev": "ry"})
    assert added.status_code == 201 and added.json()["faculty_abbrev"] == "RY" and added.json()["faculty_name"] == "Riya Yadav"
    hdr = client.patch(f"/api/v1/timetables/{tid}", headers=_h(coord), json={"division": "z8"})
    assert hdr.json()["class_key"] == "SY-Z8"


def test_permissions_and_alias_rules(client, db_session, run_sync):
    fac = register(client, "tt7-fac@test.com", role="employee")
    stu = register(client, "tt7-stu@test.com", role="student")
    coord = register(client, "tt7-coord@test.com", role="manager")
    assert _upload(client, fac, _docx("Z9")).status_code == 403
    assert client.get("/api/v1/timetables", headers=_h(fac)).status_code == 403
    res = client.post("/api/v1/timetables/aliases", headers=_h(coord),
                      json={"abbrev": "ST", "user_id": _user_id(db_session, "tt7-stu@test.com")})
    assert res.status_code == 400  # students can't be mapped
    res = client.post("/api/v1/timetables/aliases", headers=_h(coord),
                      json={"abbrev": "TOOLONG", "user_id": _user_id(db_session, "tt7-fac@test.com")})
    assert res.status_code == 422
    assert stu


def test_csv_and_xlsx_grids(client, db_session, run_sync):
    import openpyxl

    coord = _setup_people(client, "tt8")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Class: T.Y. B.TECH. Div: Z9"])
    w.writerows([["Time", "Monday", "Tuesday"], ["8:00-9:00", "OS (QX) 31", ""], ["9:00-10:00", "OS (QX) 31", "AI (RY) 32"]])
    tid = _upload(client, coord, buf.getvalue().encode(), name="z9.csv", ctype="text/csv").json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    os_ = next(e for e in d["entries"] if e["subject"] == "OS")
    assert (os_["start_time"], os_["end_time"]) == ("08:00", "10:00") and d["class_key"] == "TY-Z9"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Class: L.Y. B.TECH  Div: Z10"])
    ws.append(["Time", "Monday"])
    ws.append(["1:00-2:00", "NLP_Z101 (QZ) L7"])
    ws.append(["2:00-3:00", None])
    ws.merge_cells("B3:B4")
    out = io.BytesIO()
    wb.save(out)
    tid = _upload(client, coord, out.getvalue(), name="z10.xlsx",
                  ctype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    e = d["entries"][0]
    assert (e["subject"], e["start_time"], e["end_time"], e["kind"]) == ("NLP", "13:00", "15:00", "lab")


def test_photo_transcription_is_cross_checked(client, db_session, run_sync):
    # Initials unique to this test: aliases are global and other tests map QX/RY/QZ.
    coord = register(client, "tt9-coord@test.com", role="manager")
    register(client, "tt9-vw@test.com", full_name="Vera Wong", role="employee")
    first = {"header": {"class": "Class: T.Y. B.TECH.", "division": "Div: Z11"},
             "cells": [{"day": "Monday", "start": "8:00", "end": "9:00", "text": "ML (VW) 27"},
                       {"day": "Monday", "start": "9:00", "end": "10:00", "text": "CN (UV) 28"}],
             "legend": [{"code": "ML", "subject": "Machine Learning", "faculty": "Prof. Vera Wong"}], "unreadable": []}
    second = {**first, "cells": [first["cells"][0], {"day": "Monday", "start": "10:00", "end": "11:00", "text": "CN (RY) 28"}]}
    llm = MagicMock()
    llm.is_available = True
    llm.client.chat.completions.create.side_effect = [
        MagicMock(choices=[MagicMock(message=MagicMock(content=json.dumps(first)))]),
        MagicMock(choices=[MagicMock(message=MagicMock(content=json.dumps(second)))]),
    ]
    from PIL import Image
    img = io.BytesIO()
    Image.new("RGB", (40, 30), "white").save(img, "PNG")
    with patch("app.services.timetable_extract.llm_service", llm):
        tid = _upload(client, coord, img.getvalue(), name="photo.png", ctype="image/png").json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    assert d["extraction_method"] == "vision" and d["class_key"] == "TY-Z11"
    flags = {e["subject"]: [i["code"] for i in e["issues"]] for e in d["entries"]}
    assert "check_source" in flags["CN"] and "check_source" not in flags["ML"]
    sugg = client.get(f"/api/v1/timetables/{tid}/suggestions", headers=_h(coord)).json()
    vw = next(s for s in sugg if s["abbrev"] == "VW")
    assert vw["legend_names"] == ["Prof. Vera Wong"]
    assert [c["email"] for c in vw["candidates"]] == ["tt9-vw@test.com"]
    assert "legend" in vw["candidates"][0]["reason"] and not vw["ambiguous"]


def test_photo_without_llm_leaves_a_manual_draft(client, db_session, run_sync):
    coord = register(client, "tt10-coord@test.com", role="manager")
    llm = MagicMock()
    llm.is_available = False
    with patch("app.services.timetable_extract.llm_service", llm):
        tid = _upload(client, coord, b"\x89PNG fake", name="scan.png", ctype="image/png").json()["id"]
    d = client.get(f"/api/v1/timetables/{tid}", headers=_h(coord)).json()
    assert d["status"] == "draft" and "language model" in d["extraction_error"] and d["entries"] == []
