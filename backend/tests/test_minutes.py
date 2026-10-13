"""Meeting minutes export (DOCX / PDF), coordinator+ only."""
import io
import json
from datetime import datetime, timezone

from sqlmodel import select

from app.models.structured import AuditLog, Commitment, Employee, Meeting
from tests.helpers import register


def _h(token):
    return {"Authorization": f"Bearer {token}"}


def _meeting(db, department="all"):
    m = Meeting(
        title="Lab Readiness Sync",
        occurred_at=datetime(2026, 10, 8, tzinfo=timezone.utc),
        summary="Reviewed lab readiness ahead of practical exams.",
        participants=json.dumps(["Dr. Rohan Mehta", "Prof. Pooja Verma"]),
        decisions=json.dumps(["Upgrade the Lab 4 switch before mid-terms"]),
        department=department,
    )
    db.add(m)
    db.commit()
    emp = Employee(display_name="Prof. Pooja Verma")
    db.add(emp)
    db.commit()
    db.add(Commitment(
        assignee_employee_id=emp.id, description="Replace the Lab 4 network switch",
        source_meeting_id=m.id, due_date=datetime(2026, 10, 15, tzinfo=timezone.utc),
    ))
    db.commit()
    return m


class TestBuild:
    def test_build_collects_decisions_and_actions(self, db_session):
        from app.services.minutes_service import minutes_service

        m = minutes_service.build(db_session, _meeting(db_session))
        assert m["decisions"] == ["Upgrade the Lab 4 switch before mid-terms"]
        assert m["actions"][0] == {
            "description": "Replace the Lab 4 network switch", "owner": "Prof. Pooja Verma",
            "due": "2026-10-15", "status": "open",
        }
        assert m["attendees"] == ["Dr. Rohan Mehta", "Prof. Pooja Verma"]
        assert "not an approved record" in m["notice"]

    def test_docx_contains_the_record(self, db_session):
        from docx import Document

        from app.services.minutes_service import minutes_service

        data = minutes_service.to_docx(minutes_service.build(db_session, _meeting(db_session)))
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Minutes: Lab Readiness Sync" in text and "Upgrade the Lab 4 switch" in text
        assert doc.tables[0].rows[1].cells[1].text == "Prof. Pooja Verma"

    def test_pdf_renders(self, db_session):
        from app.services.minutes_service import minutes_service

        data = minutes_service.to_pdf(minutes_service.build(db_session, _meeting(db_session)))
        assert data.startswith(b"%PDF") and len(data) > 1000


class TestEndpoint:
    def test_coordinator_downloads_docx_and_it_is_audited(self, client, db_session):
        m = _meeting(db_session)
        token = register(client, "minutes-coord@test.com", role="manager")
        res = client.get(f"/api/v1/transcripts/{m.id}/minutes?format=docx", headers=_h(token))
        assert res.status_code == 200
        assert "wordprocessingml" in res.headers["content-type"]
        assert 'filename="Minutes_2026-10-08_Lab_Readiness_Sync.docx"' in res.headers["content-disposition"]
        audit = db_session.exec(select(AuditLog).where(AuditLog.action == "minutes_export")).all()
        assert any(a.target_id == str(m.id) for a in audit)

    def test_pdf_format(self, client, db_session):
        m = _meeting(db_session)
        token = register(client, "minutes-coord2@test.com", role="admin")
        res = client.get(f"/api/v1/transcripts/{m.id}/minutes?format=pdf", headers=_h(token))
        assert res.status_code == 200 and res.content.startswith(b"%PDF")

    def test_faculty_and_students_cannot_export(self, client, db_session):
        m = _meeting(db_session)
        for role in ("employee", "student"):
            token = register(client, f"minutes-{role}@test.com", role=role)
            assert client.get(f"/api/v1/transcripts/{m.id}/minutes", headers=_h(token)).status_code == 403

    def test_bad_format_rejected(self, client, db_session):
        m = _meeting(db_session)
        token = register(client, "minutes-coord3@test.com", role="manager")
        assert client.get(f"/api/v1/transcripts/{m.id}/minutes?format=xls", headers=_h(token)).status_code == 422
